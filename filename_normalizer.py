import os
import re
import shutil
import subprocess
from datetime import datetime
from docx import Document

# 16개 사업소 표준 매핑 정보 (광명이 제외되면서 13~16번으로 한 칸씩 당겨짐)
# - num: 파일명에 부여될 표준 번호 ("01" ~ "16")
# - idx: 마스터 템플릿(00번 워드 파일) 내에서의 테이블 인덱스 (0 ~ 16, 광명 자리인 12번 인덱스는 건너뜀)
BRANCH_DEFINITIONS = [
    {
        "num": "01",
        "idx": 0,
        "std_name_weekly": "T-타워 사업부",
        "std_name_monthly": "T-타워 사업부",
        "keywords": ["t-타워", "t타워", "티타워"],
    },
    {
        "num": "02",
        "idx": 1,
        "std_name_weekly": "수2",
        "std_name_monthly": "수2",
        "keywords": ["수도권2", "수도권 2", "수2", "삼화2호", "삼화 2호"],
    },
    {
        "num": "03",
        "idx": 2,
        "std_name_weekly": "중부",
        "std_name_monthly": "중부",
        "keywords": ["중부"],
    },
    {
        "num": "04",
        "idx": 3,
        "std_name_weekly": "부산",
        "std_name_monthly": "부산",
        "keywords": ["부산"],
    },
    {
        "num": "05",
        "idx": 4,
        "std_name_weekly": "대구",
        "std_name_monthly": "대구",
        "keywords": ["대구"],
    },
    {
        "num": "06",
        "idx": 5,
        "std_name_weekly": "서부",
        "std_name_monthly": "서부",
        "keywords": ["서부"],
    },
    {
        "num": "07",
        "idx": 6,
        "std_name_weekly": "제주",
        "std_name_monthly": "제주",
        "keywords": ["제주"],
    },
    {
        "num": "08",
        "idx": 7,
        "std_name_weekly": "서린",
        "std_name_monthly": "서린",
        "keywords": ["서린"],
    },
    {
        "num": "09",
        "idx": 8,
        "std_name_weekly": "판교",
        "std_name_monthly": "판교",
        "keywords": ["sk-p", "skp", "판교사옥", "판교"], # 주의: '판교테크원'을 먼저 체크해야 함!
    },
    {
        "num": "10",
        "idx": 9,
        "std_name_weekly": "상암 ICN10",
        "std_name_monthly": "상암 ICN10",
        "keywords": ["상암", "icn10", "icn 10", "디지털리얼티", "디지털 리얼티"],
    },
    {
        "num": "11",
        "idx": 10,
        "std_name_weekly": "종로타워",
        "std_name_monthly": "종로타워",
        "keywords": ["종로타워", "종로"],
    },
    {
        "num": "12",
        "idx": 11,
        "std_name_weekly": "삼화타워",
        "std_name_monthly": "삼화타워",
        "keywords": ["삼화타워", "삼화"], # 주의: '삼화2호'(02번)를 먼저 체크해야 함!
    },
    # (구 13번 광명 제외 -> 뒷번호 13~16번으로 한 칸씩 당김)
    {
        "num": "13",
        "idx": 13,
        "std_name_weekly": "KT&G세종",
        "std_name_monthly": "KT&G세종공장",
        "keywords": ["kt&g", "ktg", "세종공장", "세종"],
    },
    {
        "num": "14",
        "idx": 14,
        "std_name_weekly": "판교테크원",
        "std_name_monthly": "판교 테크원",
        "keywords": ["판교테크원", "판교 테크원", "테크원"],
    },
    {
        "num": "15",
        "idx": 15,
        "std_name_weekly": "SK무의연수원",
        "std_name_monthly": "SK무의연수원",
        "keywords": ["무의연수원", "sk무의", "무의"],
    },
    {
        "num": "16",
        "idx": 16,
        "std_name_weekly": "인재개발원",
        "std_name_monthly": "인재개발원",
        "keywords": ["인재개발원", "인재 개발원"],
    },
]

# 우선순위 매칭 순서 (더 구체적인 키워드를 가진 사업소를 먼저 검사: 14 판교테크원 -> 09 판교, 02 삼화2호 -> 12 삼화타워 등)
MATCH_PRIORITY_ORDER = [
    "14", # 판교테크원 ('판교'보다 먼저)
    "02", # 수2 / 삼화2호 ('삼화'보다 먼저)
    "01", # T-타워
    "10", # 상암 ICN10
    "11", # 종로타워
    "12", # 삼화타워
    "13", # KT&G세종
    "15", # SK무의연수원
    "16", # 인재개발원
    "03", # 중부
    "04", # 부산
    "05", # 대구
    "06", # 서부
    "07", # 제주
    "08", # 서린
    "09", # 판교 (SK-P타워)
]

BRANCH_BY_NUM = {b["num"]: b for b in BRANCH_DEFINITIONS}
BRANCH_BY_IDX = {b["idx"]: b for b in BRANCH_DEFINITIONS}


def convert_doc_to_docx(doc_path, target_docx_path=None, remove_original=True):
    """
    구형 워드(.doc) 파일을 MS Word COM 자동화(win32com)를 통해 .docx 포맷으로 변환합니다.
    반환값: 변환된 .docx 파일의 절대 경로
    """
    import pythoncom
    import win32com.client

    abs_doc = os.path.abspath(doc_path)
    abs_docx = os.path.abspath(target_docx_path) if target_docx_path else (os.path.splitext(abs_doc)[0] + ".docx")

    # 만약 파일 헤더가 이미 PK(ZIP 기반 .docx)인데 확장자만 .doc인 경우 복사/이름변경으로 즉시 처리
    try:
        with open(abs_doc, "rb") as f:
            magic = f.read(4)
        if magic == b"PK\x03\x04":
            if abs_doc != abs_docx:
                if os.path.exists(abs_docx):
                    os.remove(abs_docx)
                if remove_original:
                    os.rename(abs_doc, abs_docx)
                else:
                    shutil.copy2(abs_doc, abs_docx)
            return abs_docx
    except Exception:
        pass

    # 잔여 Word 프로세스 정리 후 COM 초기화
    try:
        subprocess.run(['taskkill', '/F', '/IM', 'WINWORD.EXE'], capture_output=True)
    except Exception:
        pass

    pythoncom.CoInitialize()
    word = None
    try:
        if os.path.exists(abs_docx):
            try:
                os.remove(abs_docx)
            except Exception:
                pass

        word = win32com.client.DispatchEx('Word.Application')
        word.Visible = False
        word.DisplayAlerts = 0
        doc = word.Documents.Open(abs_doc, ReadOnly=True, AddToRecentFiles=False, Visible=False)
        # FileFormat=16: wdFormatXMLDocument (.docx)
        doc.SaveAs2(abs_docx, FileFormat=16)
        doc.Close(False)
        word.Quit()
        word = None

        if remove_original and os.path.exists(abs_docx) and os.path.exists(abs_doc) and abs_doc.lower() != abs_docx.lower():
            try:
                os.remove(abs_doc)
            except Exception:
                pass

        return abs_docx
    except Exception as e:
        if word:
            try:
                word.Quit()
            except Exception:
                pass
        try:
            subprocess.run(['taskkill', '/F', '/IM', 'WINWORD.EXE'], capture_output=True)
        except Exception:
            pass
        raise RuntimeError(f".doc -> .docx 변환 실패 ({os.path.basename(doc_path)}): {e}")
    finally:
        pythoncom.CoUninitialize()


def _match_branch_from_text(text):
    """주어진 텍스트에서 우선순위에 따라 사업소 정보를 찾아 반환합니다."""
    if not text:
        return None
    norm = text.lower().replace(" ", "").replace("_", "").replace("-", "")
    norm_with_hyphen = text.lower().replace(" ", "")

    # 광명 파일이 혹시 들어오면 제외 처리될 수 있도록 None 반환 전에 체크
    for gw_kw in ["광명", "u-플래닛", "u플래닛", "유타워", "u타워", "아이벡스"]:
        if gw_kw in norm_with_hyphen or gw_kw.replace("-", "") in norm:
            return None

    for num in MATCH_PRIORITY_ORDER:
        b = BRANCH_BY_NUM[num]
        for kw in b["keywords"]:
            kw_clean = kw.lower().replace(" ", "")
            kw_no_hyphen = kw_clean.replace("-", "")
            if kw_clean in norm_with_hyphen or (len(kw_no_hyphen) >= 2 and kw_no_hyphen in norm):
                return b
    return None


def extract_date_from_filename_or_doc(filename, filepath=None):
    """파일명이나 문서 내용에서 날짜(YYYYMMDD)를 추출합니다."""
    base = os.path.splitext(os.path.basename(filename))[0]
    
    # 1. 20261023 또는 2026.10.23 또는 2026-10-23 형태
    m = re.search(r'(20\d{2})[\.\-\s/]*([01]\d)[\.\-\s/]*([0-3]\d)', base)
    if m:
        return f"{m.group(1)}{m.group(2)}{m.group(3)}"
        
    # 2. 261023 또는 26.10.23 형태 (앞에 20이 생략된 6자리)
    m6 = re.search(r'(?<!\d)(2[4-9])[\.\-]?([01]\d)[\.\-]?([0-3]\d)(?!\d)', base)
    if m6:
        return f"20{m6.group(1)}{m6.group(2)}{m6.group(3)}"

    # 3. 10.22 형태 (월.일만 있는 경우 -> 현재 연도 결합)
    m_md = re.search(r'(?<!\d)(0?[1-9]|1[0-2])[\.](0?[1-9]|[12]\d|3[01])(?!\d)', base)
    if m_md:
        year = datetime.now().year
        return f"{year}{int(m_md.group(1)):02d}{int(m_md.group(2)):02d}"

    # 4. 문서 내부에서 날짜 추출 시도
    if filepath and os.path.exists(filepath) and filepath.lower().endswith('.docx'):
        try:
            doc = Document(filepath)
            if doc.tables and len(doc.tables[0].rows) > 0:
                t_text = doc.tables[0].rows[0].cells[0].text
                dm = re.search(r'(20\d{2})\s*[\.\년]\s*(\d{1,2})\s*[\.\월]\s*(\d{1,2})', t_text)
                if dm:
                    return f"{dm.group(1)}{int(dm.group(2)):02d}{int(dm.group(3)):02d}"
            for p in doc.paragraphs[:5]:
                dm = re.search(r'(20\d{2})\s*[\.\년]\s*(\d{1,2})\s*[\.\월]\s*(\d{1,2})', p.text)
                if dm:
                    return f"{dm.group(1)}{int(dm.group(2)):02d}{int(dm.group(3)):02d}"
        except Exception:
            pass

    return datetime.now().strftime("%Y%m%d")


def detect_report_type(filename, filepath=None):
    """파일명 또는 문서 구조로 주간('weekly') / 월간('monthly') 여부를 판별합니다."""
    if "월간" in filename:
        return "monthly"
    if "주간" in filename:
        return "weekly"
    if filepath and os.path.exists(filepath) and filepath.lower().endswith('.docx'):
        try:
            doc = Document(filepath)
            for p in doc.paragraphs[:5]:
                if "월간" in p.text:
                    return "monthly"
        except Exception:
            pass
    return "weekly"


def identify_branch(filename, filepath=None):
    """
    파일명과 (필요시) 워드 문서 내부의 헤더 텍스트를 분석하여
    어떤 사업소인지 정확히 판별합니다.
    """
    base = os.path.splitext(os.path.basename(filename))[0]

    # 통합본/템플릿 파일(00번 또는 '작성예정', '작성완료', '종합', '통합')은 제외
    if base.startswith("00") or any(w in base for w in ["작성예정", "작성완료", "최종결과", "통합_작성완료", "- 종합"]):
        return None

    # 1단계: 파일명에서 사업소 키워드 검색 (가장 정확)
    branch = _match_branch_from_text(base)
    if branch:
        return branch

    # 2단계: 파일명에 사업소명이 없거나 모호하면 워드 문서 내부 텍스트(표 첫 줄 / 상단 문단) 검사
    if filepath and os.path.exists(filepath) and filepath.lower().endswith('.docx'):
        try:
            doc = Document(filepath)
            if doc.tables and len(doc.tables[0].rows) > 0:
                header_text = doc.tables[0].rows[0].cells[0].text.strip()
                branch = _match_branch_from_text(header_text)
                if branch:
                    return branch
            for p in doc.paragraphs[:5]:
                branch = _match_branch_from_text(p.text.strip())
                if branch:
                    return branch
        except Exception:
            pass

    # 3단계: 그래도 못 찾았을 때만 파일명 맨 앞 2자리 숫자(01~16) 참고
    m_num = re.match(r'^(\d{2})', base.strip())
    if m_num and m_num.group(1) in BRANCH_BY_NUM:
        return BRANCH_BY_NUM[m_num.group(1)]

    return None


def build_standard_filename(branch_info, date_str, report_type="weekly", ext=".docx"):
    """표준 파일명을 생성합니다."""
    num = branch_info["num"]
    if report_type == "monthly":
        b_name = branch_info["std_name_monthly"]
        return f"{num}.월간 업무 보고(대표이사님 보고용)_{date_str}_{b_name}{ext}"
    else:
        b_name = branch_info["std_name_weekly"]
        return f"{num}.주간경영회의자료({date_str})_{b_name}{ext}"


def normalize_filename(filename, filepath=None, force_report_type=None, force_date=None, convert_doc=True):
    """
    단일 파일에 대해 사업소를 판별하고 표준 파일명을 반환합니다.
    convert_doc=True 이면 .doc 확장자도 최종 .docx 확장자로 표준 파일명을 생성합니다.
    반환값: (branch_info, new_filename, detected_date) 또는 인식 실패시 (None, filename, None)
    """
    ext = os.path.splitext(filename)[1]
    if convert_doc and ext.lower() == ".doc":
        ext = ".docx"
    branch_info = identify_branch(filename, filepath)
    if not branch_info:
        return None, filename, None

    rtype = force_report_type or detect_report_type(filename, filepath)
    date_str = force_date or extract_date_from_filename_or_doc(filename, filepath)
    new_name = build_standard_filename(branch_info, date_str, rtype, ext)
    return branch_info, new_name, date_str


def batch_rename_in_folder(folder_path, unify_date=True, target_date=None, dry_run=False, convert_doc=True):
    """
    지정된 폴더 내의 제각각인 보고서 파일들을 스캔하여 표준 파일명(01~16번)으로 일괄 변경합니다.
    convert_doc=True(기본값)이면 삼화타워, 판교 등의 구형 .doc 파일을 자동으로 .docx로 변환합니다.
    """
    if not os.path.isdir(folder_path):
        raise ValueError(f"폴더를 찾을 수 없습니다: {folder_path}")

    all_files = sorted(os.listdir(folder_path))
    all_files_set = set(all_files)

    candidates = []
    for fname in all_files:
        if fname.startswith('~') or fname.startswith('.'):
            continue
        base_f, ext = os.path.splitext(fname)
        ext_low = ext.lower()
        if ext_low not in ['.docx', '.doc']:
            continue
        fpath = os.path.join(folder_path, fname)
        if not os.path.isfile(fpath):
            continue
        
        if fname.startswith("00") or any(w in fname for w in ["작성예정", "작성완료", "최종결과", "통합_작성완료", "- 종합"]):
            continue

        # 만약 동일한 이름의 .doc와 .docx가 이미 폴더에 같이 들어있다면, 중복 방지를 위해 .doc는 건너뛰거나 정리
        if ext_low == '.doc' and (base_f + '.docx') in all_files_set:
            continue

        b_info = identify_branch(fname, fpath)
        rtype = detect_report_type(fname, fpath)
        d_str = extract_date_from_filename_or_doc(fname, fpath)
        candidates.append({
            "old_name": fname,
            "old_path": fpath,
            "ext": ext,
            "is_doc": ext_low == ".doc",
            "branch": b_info,
            "report_type": rtype,
            "date": d_str,
        })

    common_date_by_type = {}
    if target_date and re.match(r'^\d{8}$', target_date.strip()):
        common_date_by_type["weekly"] = target_date.strip()
        common_date_by_type["monthly"] = target_date.strip()
    elif unify_date and candidates:
        for rt in ["weekly", "monthly"]:
            dates = [c["date"] for c in candidates if c["report_type"] == rt and c["date"]]
            if dates:
                counts = {}
                for d in dates:
                    counts[d] = counts.get(d, 0) + 1
                best_date = sorted(counts.keys(), key=lambda x: (counts[x], x), reverse=True)[0]
                common_date_by_type[rt] = best_date

    results = []
    used_target_names = set()
    temp_moves = []

    for item in candidates:
        b_info = item["branch"]
        is_doc = item["is_doc"]

        # 만약 파일명에서 사업소가 안 잡혔는데 .doc 파일이고 실제 실행(dry_run=False)이라면, 먼저 .docx로 변환해서 내부 텍스트로 사업소 재탐색
        if not b_info and is_doc and convert_doc and not dry_run:
            try:
                converted_tmp = convert_doc_to_docx(item["old_path"], remove_original=True)
                item["old_path"] = converted_tmp
                item["ext"] = ".docx"
                b_info = identify_branch(os.path.basename(converted_tmp), converted_tmp)
                item["branch"] = b_info
            except Exception:
                pass

        if not b_info:
            results.append({
                "old_name": item["old_name"],
                "new_name": item["old_name"],
                "branch_num": "제외/미인식",
                "branch_name": "광명 또는 미인식",
                "status": "skipped",
                "message": "제외 대상(광명) 또는 사업소 미인식",
            })
            continue

        rt = item["report_type"]
        final_date = common_date_by_type.get(rt, item["date"])
        target_ext = ".docx" if (convert_doc and is_doc) else item["ext"]
        new_name = build_standard_filename(b_info, final_date, rt, target_ext)

        if new_name in used_target_names:
            base_n, ext_n = os.path.splitext(new_name)
            dup_idx = 2
            while f"{base_n}_{dup_idx}{ext_n}" in used_target_names:
                dup_idx += 1
            new_name = f"{base_n}_{dup_idx}{ext_n}"
        used_target_names.add(new_name)

        if item["old_name"] == new_name:
            results.append({
                "old_name": item["old_name"],
                "new_name": new_name,
                "branch_num": b_info["num"],
                "branch_name": b_info["std_name_weekly"],
                "status": "unchanged",
                "message": "이미 표준 파일명입니다.",
            })
            continue

        if dry_run:
            msg = "doc→docx 변환 및 이름 변경 예정" if (convert_doc and is_doc) else "변경 예정"
            results.append({
                "old_name": item["old_name"],
                "new_name": new_name,
                "branch_num": b_info["num"],
                "branch_name": b_info["std_name_weekly"],
                "status": "preview",
                "message": msg,
            })
        else:
            temp_name = f"__renaming_tmp_{b_info['num']}_{id(item)}{target_ext}"
            temp_path = os.path.join(folder_path, temp_name)
            try:
                if convert_doc and is_doc and item["old_path"].lower().endswith(".doc"):
                    convert_doc_to_docx(item["old_path"], target_docx_path=temp_path, remove_original=True)
                else:
                    os.rename(item["old_path"], temp_path)
                temp_moves.append((temp_path, os.path.join(folder_path, new_name), item, new_name, is_doc and convert_doc))
            except Exception as e:
                results.append({
                    "old_name": item["old_name"],
                    "new_name": new_name,
                    "branch_num": b_info["num"],
                    "branch_name": b_info["std_name_weekly"],
                    "status": "error",
                    "message": f"변환/변경 실패: {str(e)}",
                })

    if not dry_run:
        for temp_path, final_path, item, new_name, was_converted in temp_moves:
            b_info = item["branch"]
            try:
                if os.path.exists(final_path):
                    base_f, ext_f = os.path.splitext(final_path)
                    backup_path = f"{base_f}_기존파일백업{ext_f}"
                    shutil.move(final_path, backup_path)
                os.rename(temp_path, final_path)
                msg = "doc→docx 변환 및 변경 완료" if was_converted else "변경 완료"
                results.append({
                    "old_name": item["old_name"],
                    "new_name": new_name,
                    "branch_num": b_info["num"],
                    "branch_name": b_info["std_name_weekly"],
                    "status": "renamed",
                    "message": msg,
                })
            except Exception as e:
                try:
                    os.rename(temp_path, item["old_path"])
                except Exception:
                    pass
                results.append({
                    "old_name": item["old_name"],
                    "new_name": new_name,
                    "branch_num": b_info["num"],
                    "branch_name": b_info["std_name_weekly"],
                    "status": "error",
                    "message": f"최종 이름 변경 실패: {str(e)}",
                })

    results.sort(key=lambda r: r["branch_num"])
    return results
