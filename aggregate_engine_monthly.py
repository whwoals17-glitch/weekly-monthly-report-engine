from docx.oxml.ns import qn
import os
import copy
from docx import Document
from docx.shared import Pt, Cm

from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH

from aggregate_engine import force_font_on_cell

def write_exact_text_to_cell(cell, text, font_name="가는각진제목체", font_size_pt=10, is_bold=False, alignment=None, space_pt=3):
    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import Pt, RGBColor
    
    p = cell.paragraphs[0]
    p.text = ""
    pPr = p._element.get_or_add_pPr()
    for tag in ['w:numPr', 'w:ind', 'w:tabs', 'w:spacing', 'w:pStyle']:
        elem = pPr.find(qn(tag))
        if elem is not None:
            pPr.remove(elem)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_before = Pt(space_pt)
    p.paragraph_format.space_after = Pt(space_pt)
    if alignment is not None:
        p.alignment = alignment
    for extra_p in cell.paragraphs[1:]:
        extra_p._element.getparent().remove(extra_p._element)
    lines = text.split('\n')
    for i, line in enumerate(lines):
        if i > 0:
            p.add_run('\n')
        run = p.add_run(line)
        run.font.name = font_name
        run.font.size = Pt(font_size_pt)
        run.font.bold = is_bold
        run.font.color.rgb = RGBColor(0, 0, 0)
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.get_or_add_rFonts()
        rFonts.set(qn('w:eastAsia'), font_name)

def copy_cell_formatting(source_cell, target_cell):
    source_tcPr = source_cell._tc.get_or_add_tcPr()
    target_tcPr = target_cell._tc.get_or_add_tcPr()
    
    source_shd = source_tcPr.find(qn('w:shd'))
    if source_shd is not None:
        target_shd = target_tcPr.find(qn('w:shd'))
        if target_shd is not None:
            target_tcPr.remove(target_shd)
        target_tcPr.append(copy.deepcopy(source_shd))
        
    source_vAlign = source_tcPr.find(qn('w:vAlign'))
    if source_vAlign is not None:
        target_vAlign = target_tcPr.find(qn('w:vAlign'))
        if target_vAlign is not None:
            target_tcPr.remove(target_vAlign)
        target_tcPr.append(copy.deepcopy(source_vAlign))

def set_row_height(row, height_cm):
    trPr = row._tr.get_or_add_trPr()
    trHeight = trPr.find(qn('w:trHeight'))
    if trHeight is None:
        trHeight = OxmlElement('w:trHeight')
        trPr.append(trHeight)
    # 1 cm = 567 twips
    trHeight.set(qn('w:val'), str(int(height_cm * 567.0)))
    trHeight.set(qn('w:hRule'), 'atLeast')

def clean_text(text):
    return text.replace('\n', '').replace(' ', '').strip()

def process_and_merge_monthly(template_path, upload_folder, output_path):
    print("=== 월간 경영회의 자동 취합 엔진 시작 ===")
    
    master_doc = Document(template_path)
    print(f"템플릿 로드 완료: 총 {len(master_doc.tables)}개의 표 발견됨 (예상: 34개)")

    # --- Dynamic Header Update from T-타워 ---
    t_tower_file = None
    for f in os.listdir(upload_folder):
        if ('T-타워' in f or f.startswith('01')) and not f.startswith('~'):
            t_tower_file = os.path.join(upload_folder, f)
            break
            
    if t_tower_file:
        try:
            t_doc = Document(t_tower_file)
            new_title = ''
            new_date = ''
            import re
            for p in t_doc.paragraphs:
                if '월간 업무 보고' in p.text:
                    raw_title = p.text.strip()
                    m_match = re.search(r'(\d+)월', raw_title)
                    if m_match:
                        month_str = str(int(m_match.group(1)))
                        new_title = f'월간 업무 보고 ({month_str}월)'
                    else:
                        new_title = raw_title
                if '보고일자' in p.text:
                    match = re.search(r'(보고일자.*)', p.text)
                    if match:
                        new_date = match.group(1)
            
            new_col1 = ''
            new_col2 = ''
            for tbl in t_doc.tables:
                if len(tbl.rows) > 0 and len(tbl.rows[0].cells) >= 3:
                    if '중점추진' in tbl.rows[0].cells[1].text:
                        new_col1 = tbl.rows[0].cells[1].text.strip()
                        new_col2 = tbl.rows[0].cells[2].text.strip()
                        break
                        
            if new_title and new_date:
                print(f"T-타워 기준 동적 헤더 업데이트: {new_title} / {new_date}")
                
                old_title = ''
                old_date = ''
                for p in master_doc.paragraphs:
                    if '월간 업무 보고' in p.text:
                        old_title = p.text.strip()
                    if '보고일자' in p.text:
                        match = re.search(r'(보고일자.*)', p.text)
                        if match:
                            old_date = match.group(1)
                            break
                            
                old_col1 = ''
                old_col2 = ''
                for i in range(1, len(master_doc.tables), 2):
                    tbl = master_doc.tables[i]
                    if len(tbl.rows) > 0 and len(tbl.rows[0].cells) >= 3:
                        if '중점추진' in tbl.rows[0].cells[1].text:
                            old_col1 = tbl.rows[0].cells[1].text.strip()
                            old_col2 = tbl.rows[0].cells[2].text.strip()
                            break
                            
                def replace_paragraph_text(p, old_t, new_t):
                    if not p.runs or not old_t: return
                    if old_t not in p.text: return
                    
                    for r in p.runs:
                        if old_t in r.text:
                            r.text = r.text.replace(old_t, new_t)
                            return
                            
                    text_str = ''
                    run_mapping = []
                    for r_idx, r in enumerate(p.runs):
                        for char in r.text:
                            text_str += char
                            run_mapping.append(r_idx)
                            
                    start_idx = text_str.find(old_t)
                    if start_idx == -1: return
                    end_idx = start_idx + len(old_t) - 1
                    
                    start_run = run_mapping[start_idx]
                    end_run = run_mapping[end_idx]
                    
                    combined = ''
                    for i in range(start_run, end_run + 1):
                        combined += p.runs[i].text
                        
                    combined = combined.replace(old_t, new_t, 1)
                    
                    p.runs[start_run].text = combined
                    for i in range(start_run + 1, end_run + 1):
                        p.runs[i].text = ''

                for p in master_doc.paragraphs:
                    if old_title and old_title in p.text:
                        replace_paragraph_text(p, old_title, new_title)
                    if old_date and old_date in p.text:
                        replace_paragraph_text(p, old_date, new_date)

                for tbl in master_doc.tables:
                    for r in tbl.rows:
                        for c in r.cells:
                            for p in c.paragraphs:
                                if old_col1 and old_col1 in p.text:
                                    replace_paragraph_text(p, old_col1, new_col1)
                                if old_col2 and old_col2 in p.text:
                                    replace_paragraph_text(p, old_col2, new_col2)
        except Exception as e:
            print(f"동적 헤더 업데이트 중 오류 발생 (무시됨): {e}")
    # -----------------------------------------
    
    num_branches = len(master_doc.tables) // 2
    elements_to_delete = []

    for idx in range(num_branches):
        branch_num_str = f"{idx + 1:02d}"
        
        master_tables = [master_doc.tables[idx * 2], master_doc.tables[idx * 2 + 1]]
        
        t_title_elem = master_tables[0]._element.getprevious().getprevious()
        t_title = ''.join(t_title_elem.itertext()) if t_title_elem is not None else ""
        
        is_gwangmyeong = '광명' in t_title
        
        target_file = None
        for f in os.listdir(upload_folder):
            if f.startswith(branch_num_str) and not f.startswith('~'):
                target_file = os.path.join(upload_folder, f)
                break
                
        if is_gwangmyeong or not target_file:
            if not target_file:
                print(f"[{branch_num_str} 호] 파일 없음. 양식 삭제.")
            else:
                print(f"[{branch_num_str} 호] 광명 지점 제외 요청에 따라 양식 삭제.")
            
            elements_to_delete.extend([
                master_tables[0]._element.getprevious().getprevious().getprevious(),
                master_tables[0]._element.getprevious().getprevious(),
                master_tables[0]._element.getprevious(),
                master_tables[0]._element,
                master_tables[1]._element.getprevious(),
                master_tables[1]._element
            ])
            continue
            
        print(f"[{branch_num_str} 호] 파일 취합: {os.path.basename(target_file)}")
        branch_doc = Document(target_file)
        
        if len(branch_doc.tables) < 2:
            print(f"  [경고] {branch_num_str} 호 파일에 테이블이 2개 미만입니다.")
            continue
            
        branch_tables = [branch_doc.tables[0], branch_doc.tables[1]]
        master_tables = [master_doc.tables[idx * 2], master_doc.tables[idx * 2 + 1]]
        
        for t_idx in range(2):
            master_tbl = master_tables[t_idx]
            branch_tbl = branch_tables[t_idx]
            
            if t_idx == 0:
                # 다음 사업장 시작 시(Table 0 직전의 '월간 업무 보고' 단락) 페이지 나누기 속성 부여하여 이중 빈페이지 방지
                if idx > 0:
                    prev = master_tbl._element.getprevious()
                    while prev is not None:
                        if prev.tag == qn('w:p'):
                            text = "".join(node.text for node in prev.iter() if node.tag == qn('w:t') and node.text)
                            if "월간업무보고" in text.replace(" ", ""):
                                pPr = prev.get_or_add_pPr()
                                if pPr.find(qn('w:pageBreakBefore')) is None:
                                    pb = OxmlElement('w:pageBreakBefore')
                                    pPr.append(pb)
                                # 월간업무보고 앞에 글머리기호가 생기지 않도록 강제 제거
                                numPr = pPr.find(qn('w:numPr'))
                                if numPr is not None:
                                    pPr.remove(numPr)
                                pStyle = pPr.find(qn('w:pStyle'))
                                if pStyle is not None:
                                    pPr.remove(pStyle)
                                break
                        prev = prev.getprevious()
                            
                # 1. 인원현황: 틀을 유지하고 데이터만 매핑해서 덮어씀
                # 헤더 추출
                m_header_idx = 0
                for r_i in range(min(3, len(master_tbl.rows))):
                    r_texts = [clean_text(c.text) for c in master_tbl.rows[r_i].cells]
                    if "관리" in r_texts or "시설" in r_texts or "사무지원" in r_texts:
                        m_header_idx = r_i
                        break
                        
                master_headers_raw = [c.text for c in master_tbl.rows[m_header_idx].cells]
                master_headers = [clean_text(h) for h in master_headers_raw]
                
                b_header_idx = 0
                for r_i in range(min(3, len(branch_tbl.rows))):
                    r_texts = [clean_text(c.text) for c in branch_tbl.rows[r_i].cells]
                    if "관리" in r_texts or "시설" in r_texts or "사무지원" in r_texts:
                        b_header_idx = r_i
                        break
                        
                branch_headers = [clean_text(c.text) for c in branch_tbl.rows[b_header_idx].cells]

                # 행 높이 세팅 (헤더: 0.6cm, 데이터: 0.37cm)
                for i in range(m_header_idx + 1):
                    set_row_height(master_tbl.rows[i], 0.6)
                for r in range(1, 4):
                    m_r_idx = m_header_idx + r
                    if m_r_idx < len(master_tbl.rows):
                        set_row_height(master_tbl.rows[m_r_idx], 0.37)

                seen_cells = set()
                for r_idx in range(1, 4):
                    m_r_idx = m_header_idx + r_idx
                    b_r_idx = b_header_idx + r_idx
                    if m_r_idx >= len(master_tbl.rows) or b_r_idx >= len(branch_tbl.rows):
                        break

                    master_row = master_tbl.rows[m_r_idx]
                    branch_row = branch_tbl.rows[b_r_idx]
                    
                    for c_idx, m_cell in enumerate(master_row.cells):
                        if m_cell in seen_cells:
                            continue
                        seen_cells.add(m_cell)
                        
                        m_header = master_headers[c_idx]
                        
                        # 값 찾기
                        val = ""
                        if c_idx == 0:
                            # '구분' (정원, 현원, 증감)
                            val = master_row.cells[0].text.strip() # 기본값 유지
                            if branch_row.cells:
                                val = branch_row.cells[0].text.strip()
                        elif c_idx == len(master_row.cells) - 1:
                            # '비고' 열
                            try:
                                b_idx = branch_headers.index(m_header)
                                val = branch_row.cells[b_idx].text.strip()
                            except ValueError:
                                # '비고'라는 이름이 없으면 맨 마지막 열 시도
                                val = branch_row.cells[-1].text.strip()
                            if not val or val == '-' or val == '0':
                                val = '1. 결원: \n2. 대책: '
                            else:
                                import re
                                def clean_prefix(t, kw):
                                    t = t.strip()
                                    t = re.sub(r'^(?:[-*•\s]*)(?:(?:\d+|[가-힣])\s*\.)?\s*', '', t).strip()
                                    if kw == '결원':
                                        t = re.sub(r'^결\s*원\s*[:：]?\s*', '', t, flags=re.IGNORECASE).strip()
                                    elif kw == '대책':
                                        t = re.sub(r'^대\s*책\s*[:：]?\s*', '', t, flags=re.IGNORECASE).strip()
                                    t = re.sub(r'^(?:[-*•\s]*)(?:(?:\d+|[가-힣])\s*\.)?\s*', '', t).strip()
                                    t = re.sub(r'^[:：]\s*', '', t).strip()
                                    return t
                                    
                                match_d = re.search(r'(?:^|\n)(?:[-*•\d\.\s가-힣]*)(대\s*책)[:：]?\s*(.*)', val, re.IGNORECASE | re.DOTALL)
                                if match_d:
                                    daechaek = match_d.group(2).strip()
                                    gyulwon = val[:match_d.start()].strip()
                                else:
                                    match_12 = re.search(r'(?:^|\n)(?:[-*•\s]*)(?:1|가)\.\s*(.*?)\s*(?:^|\n)(?:[-*•\s]*)(?:2|나)\.\s*(.*)', val, re.IGNORECASE | re.DOTALL)
                                    if match_12:
                                        gyulwon = match_12.group(1).strip()
                                        daechaek = match_12.group(2).strip()
                                    else:
                                        match_job = re.search(r'([\(,\-\s]*)((?:채용|구인|면접).*)', val, re.IGNORECASE | re.DOTALL)
                                        if match_job:
                                            gyulwon = val[:match_job.start()].strip()
                                            daechaek = match_job.group(2).strip()
                                            daechaek = re.sub(r'[\)\.]+$', '', daechaek).strip()
                                            gyulwon = re.sub(r'[\(\-,\s]+$', '', gyulwon).strip()
                                        else:
                                            gyulwon = val
                                            daechaek = ''
                                            
                                gyulwon = clean_prefix(gyulwon, '결원')
                                daechaek = clean_prefix(daechaek, '대책')
                                
                                out = []
                                out.append(f"1. 결원: {gyulwon}")
                                out.append(f"2. 대책: {daechaek}")
                                val = '\n'.join(out)
                        else:
                            # 중간의 부서 열들
                            try:
                                b_idx = branch_headers.index(m_header)
                                val = branch_row.cells[b_idx].text.strip()
                                # 빈칸이면 0으로
                                if not val or val == '-':
                                    val = '0'
                            except ValueError:
                                if m_header == "보안(서린)":
                                    try:
                                        b_idx = branch_headers.index("보안")
                                        val = branch_row.cells[b_idx].text.strip()
                                        if not val or val == '-':
                                            val = '0'
                                    except ValueError:
                                        val = '0'
                                elif m_header == "CS":
                                    try:
                                        b_idx = branch_headers.index("안내")
                                        val = branch_row.cells[b_idx].text.strip()
                                        if not val or val == '-':
                                            val = '0'
                                    except ValueError:
                                        val = '0'
                                else:
                                    val = '0' # 못 찾으면 0 처리
                            if "보안" in m_header:
                                try:
                                    lg_idx = branch_headers.index("라이프가드")
                                    lg_val = branch_row.cells[lg_idx].text.strip()
                                    if lg_val and lg_val != '-':
                                        v1 = int(val) if val.replace('-', '').isdigit() else 0
                                        v2 = int(lg_val) if lg_val.replace('-', '').isdigit() else 0
                                        val = str(v1 + v2)
                                except ValueError:
                                    pass
                                
                        # 정렬: 비고(마지막 열)은 LEFT, 나머지는 CENTER
                        align = WD_ALIGN_PARAGRAPH.LEFT if c_idx == len(master_row.cells) - 1 else WD_ALIGN_PARAGRAPH.CENTER
                        force_font_on_cell(m_cell, val, font_name="가는각진제목체", font_size_pt=10, alignment=align, space_pt=0)

                # Table 0 전체에 대해 '비고' 열(마지막 열)을 제외하고 단락 줄간격 배수 1.15 강제 설정
                for row_idx in range(len(master_tbl.rows)):
                    t_row = master_tbl.rows[row_idx]
                    for col_idx, t_cell in enumerate(t_row.cells):
                        if col_idx != len(t_row.cells) - 1:
                            for p in t_cell.paragraphs:
                                p.paragraph_format.line_spacing = 1.15

            else:
                # 2. 실적 및 계획: 00템플릿의 8개 고정 구분(기획업무 등)을 그대로 유지하고 데이터만 매핑
                
                # 헤더(0번 행) 강제 통일 및 10pt 고정 (굵게 해제)
                t_tower_header_1 = master_doc.tables[1].rows[0].cells[1].text.strip()
                t_tower_header_2 = master_doc.tables[1].rows[0].cells[2].text.strip()
                
                write_exact_text_to_cell(master_tbl.rows[0].cells[0], "구분", font_name="가는각진제목체", font_size_pt=10, is_bold=False, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_pt=3)
                write_exact_text_to_cell(master_tbl.rows[0].cells[1], t_tower_header_1, font_name="가는각진제목체", font_size_pt=10, is_bold=False, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_pt=3)
                write_exact_text_to_cell(master_tbl.rows[0].cells[2], t_tower_header_2, font_name="가는각진제목체", font_size_pt=10, is_bold=False, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_pt=3)

                # branch_tbl에서 카테고리별 데이터 추출
                branch_data = {}
                def extract_cell_text_with_bullets(cell):
                    lines = []
                    for p in cell.paragraphs:
                        text = p.text
                        pPr = p._element.pPr
                        if pPr is not None and pPr.find(qn('w:numPr')) is not None:
                            text = "· " + text
                        lines.append(text)
                    return '\n'.join(lines).strip()

                last_cat_clean = None
                for r_idx, branch_row in enumerate(branch_tbl.rows):
                    if r_idx == 0:
                        continue
                    cat_raw = branch_row.cells[0].text
                    cat_clean = clean_text(cat_raw)
                    if cat_clean == "고객사주요동향":
                        cat_clean = "고객사동향"

                    if not cat_clean and last_cat_clean:
                        cat_clean = last_cat_clean
                    elif not cat_clean:
                        continue
                        
                    last_cat_clean = cat_clean
                    
                    col1 = extract_cell_text_with_bullets(branch_row.cells[1]) if len(branch_row.cells) > 1 else ""
                    col2 = extract_cell_text_with_bullets(branch_row.cells[2]) if len(branch_row.cells) > 2 else ""
                    
                    if cat_clean in branch_data:
                        old_col1, old_col2 = branch_data[cat_clean]
                        new_col1 = (old_col1 + "\n" + col1).strip() if col1 else old_col1
                        new_col2 = (old_col2 + "\n" + col2).strip() if col2 else old_col2
                        branch_data[cat_clean] = (new_col1, new_col2)
                    else:
                        branch_data[cat_clean] = (col1, col2)
                
                # master_tbl의 행들을 순회하면서 branch_data 매핑 (템플릿의 고정 구분 유지)
                for r_idx in range(1, len(master_tbl.rows)):
                    master_row = master_tbl.rows[r_idx]
                    cat_raw = master_row.cells[0].text.strip()
                    cat_clean = clean_text(cat_raw)
                    
                    val1, val2 = branch_data.get(cat_clean, ("", ""))
                    
                    # 주변동종업계동향만 9pt, 나머지는 10pt
                    row_font_size = 9 if cat_clean == "주변동종업계동향" else 10
                    
                    # c_idx == 0 (구분 열) 다시 쓰기
                    write_exact_text_to_cell(master_row.cells[0], cat_raw, font_name="가는각진제목체", font_size_pt=row_font_size, is_bold=False, alignment=WD_ALIGN_PARAGRAPH.CENTER, space_pt=3)
                    
                    # c1, c2 정렬 및 쓰기 (기존 3단 로직)
                    for c_idx, text in enumerate([val1, val2], start=1):
                        target_cell = master_row.cells[c_idx]
                        align = WD_ALIGN_PARAGRAPH.LEFT
                        
                        # 지능형 3단 내어쓰기(들여쓰기) 계층 정렬 로직
                        lines = text.split('\n')
                        aligned_lines = []
                        import re
                        import unicodedata
                        
                        def get_exact_indent_string(s):
                            out = ''
                            for c in s:
                                if c == ' ':
                                    out += ' '
                                elif c == '\t':
                                    out += '    '
                                else:
                                    if unicodedata.east_asian_width(c) in ['W', 'F']:
                                        out += '\u3000'
                                    else:
                                        out += '\u2002'
                            return out

                        current_bracket_indent = ""
                        current_item_indent = ""
                        
                        for line in lines:
                            line = line.replace('\t', '  ')
                            stripped = line.lstrip()
                            if not stripped:
                                continue
                                
                            match_bracket = re.match(r'^(\[[^\]]+\]\s*)(?:([0-9]+\.)\s+)?(.*)', stripped)
                            match_num = re.match(r'^([0-9]+\.)\s*(.*)', stripped)
                            match_bullet = re.match(r'^([-※*·])\s*(.*)', stripped)
                            
                            if match_bracket:
                                bracket_part = match_bracket.group(1)
                                num_part = match_bracket.group(2)
                                rest = match_bracket.group(3)
                                
                                current_bracket_indent = get_exact_indent_string(bracket_part)
                                my_indent = ""
                                out_line = bracket_part
                                
                                if num_part:
                                    prefix = num_part + " "
                                    out_line += prefix + rest
                                    content_part = rest
                                else:
                                    out_line += rest
                                    prefix = ""
                                    content_part = rest
                                    
                                aligned_lines.append(out_line)
                                
                                colon_idx = content_part.find(':')
                                bullet_match = re.search(r'\s+([-※*·])', content_part)
                                
                                if colon_idx != -1 and colon_idx <= 15:
                                    header = prefix + content_part[:colon_idx+1]
                                    header += ' ' if len(content_part) > colon_idx + 1 and content_part[colon_idx+1] == ' ' else ' '
                                    current_item_indent = current_bracket_indent + get_exact_indent_string(header)
                                elif bullet_match and bullet_match.start() <= 15:
                                    header = prefix + content_part[:bullet_match.end()-1]
                                    current_item_indent = current_bracket_indent + get_exact_indent_string(header)
                                else:
                                    current_item_indent = current_bracket_indent + get_exact_indent_string(prefix)
                                    if not prefix:
                                        current_item_indent = current_bracket_indent + "    "
                                        
                            elif match_num:
                                num_part = match_num.group(1)
                                rest = match_num.group(2)
                                prefix = num_part + " "
                                
                                my_indent = current_bracket_indent
                                out_line = my_indent + prefix + rest
                                aligned_lines.append(out_line)
                                
                                colon_idx = rest.find(':')
                                bullet_match = re.search(r'\s+([-※*·])', rest)
                                
                                if colon_idx != -1 and colon_idx <= 15:
                                    header = prefix + rest[:colon_idx+1]
                                    header += ' ' if len(rest) > colon_idx + 1 and rest[colon_idx+1] == ' ' else ' '
                                    current_item_indent = my_indent + get_exact_indent_string(header)
                                elif bullet_match and bullet_match.start() <= 15:
                                    header = prefix + rest[:bullet_match.end()-1]
                                    current_item_indent = my_indent + get_exact_indent_string(header)
                                else:
                                    current_item_indent = my_indent + get_exact_indent_string(prefix)
                                    
                            elif match_bullet:
                                prefix = match_bullet.group(1) + " "
                                rest = match_bullet.group(2)
                                my_indent = current_item_indent
                                out_line = my_indent + prefix + rest
                                aligned_lines.append(out_line)
                                
                            else:
                                colon_idx = stripped.find(':')
                                bullet_match = re.search(r'\s+([-※*·])', stripped)
                                
                                if colon_idx != -1 and colon_idx <= 15:
                                    my_indent = current_bracket_indent
                                    out_line = my_indent + stripped
                                    aligned_lines.append(out_line)
                                    
                                    header = stripped[:colon_idx+1]
                                    header += ' ' if len(stripped) > colon_idx + 1 and stripped[colon_idx+1] == ' ' else ' '
                                    current_item_indent = my_indent + get_exact_indent_string(header)
                                elif bullet_match and bullet_match.start() <= 15:
                                    my_indent = current_bracket_indent
                                    out_line = my_indent + stripped
                                    aligned_lines.append(out_line)
                                    
                                    header = stripped[:bullet_match.end()-1]
                                    current_item_indent = my_indent + get_exact_indent_string(header)
                                else:
                                    my_indent = current_item_indent
                                    out_line = my_indent + stripped
                                    aligned_lines.append(out_line)

                        text = '\n'.join(aligned_lines)
                        
                        # c1, c2의 내용은 항상 10pt 고정 (주변 동종 업계 동향이더라도)
                        write_exact_text_to_cell(target_cell, text, font_name="가는각진제목체", font_size_pt=10, is_bold=False, alignment=align, space_pt=3)

    print("모든 데이터 취합 완료. 저장 준비 중...")
    for p in master_doc.paragraphs:
        if "판교사옥" in p.text:
            for run in p.runs:
                if "판교사옥" in run.text:
                    run.text = run.text.replace("판교사옥", "SK-P타워")
            if "판교사옥" in p.text:
                p.text = p.text.replace("판교사옥", "SK-P타워")
        # --- 페이지 나누기 정리 로직 시작 ---
    for e in elements_to_delete:
        try:
            if e is not None and e.getparent() is not None:
                e.getparent().remove(e)
        except Exception:
            pass

    first_title_found = False

    for p in master_doc.paragraphs:
        # 1. 수동 페이지 나누기(Hard Page Break) 완전 제거
        for r in p.runs:
            if 'w:br' in r._r.xml and 'type="page"' in r._r.xml:
                
                brs = r._r.findall(qn('w:br'))
                for br in brs:
                    if br.get(qn('w:type')) == 'page':
                        r._r.remove(br)
                        
        # 2. '단락 앞 페이지 나누기' 속성 부여 (첫 번째 제목 제외)
        if '월간 업무 보고' in p.text:
            if not first_title_found:
                first_title_found = True
            else:
                p.paragraph_format.page_break_before = True
    # --- 페이지 나누기 정리 로직 끝 ---
    
    master_doc.save(output_path)
    print("=== 월간 업무 보고 작업 완료 ===")
    return True
